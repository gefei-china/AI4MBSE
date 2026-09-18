"""通用报告生成服务（SP-R 升级）：多报告类型 + 自定义大纲 + 结构化数据契约 + Markdown 渲染。

- 单一数据契约：Report = {title, sections:[{heading, body, table?, chart?}], summary, report_type}
- 报告类型体系：analysis（模型分析）/ impact（变更影响分析）/ review（预评审）——各配专业分节模板
- 自定义大纲：用户输入中显式指定章节结构（如「按 背景/方案/结论 组织」）→ 动态覆盖类型模板
- 分节模板驱动：LLM 按节聚焦填充，结构由模板控制 → 格式稳定
- 双入口复用：AI 建模 report_generation 意图 与 AI 工坊报告节点均调用本服务
- 前端 renderMarkdown 统一渲染（标题/表格/图表）
"""
import json
import os
import re
import time
from typing import Optional


# 报告类型体系：类型 key → 标签/识别关键词/分节模板（SP-R）
# 分节模板对标行业工程报告规范（工程咨询 / 系统工程 / CCB 变更评估）：
# 报告信息 → 摘要 → 背景/范围/方法 → 主体分析 → 风险 → 结论 → 建议 → 附录
REPORT_TYPES = {
    "analysis": {
        "label": "工程分析报告",
        "keywords": ["分析报告", "模型分析", "工程分析", "现状分析", "整体分析", "评估报告",
                     "架构分析", "性能分析", "系统分析", "总体分析"],
        "sections": [
            {"heading": "报告信息", "desc": "文档编号、版本、日期、编制/审核、密级（用表格呈现）"},
            {"heading": "摘要", "desc": "一段话浓缩：分析背景、对象、方法与核心结论"},
            {"heading": "分析背景与目标", "desc": "问题来源、分析目的与待回答的关键问题"},
            {"heading": "分析范围与约束", "desc": "覆盖对象与边界、假设、数据来源及限制条件"},
            {"heading": "分析方法与依据", "desc": "分析方法、模型/工具、引用标准与规范"},
            {"heading": "模型结构与数据分析", "desc": "系统组成/层级/接口结构，关键指标量化对比（表格呈现）"},
            {"heading": "问题与风险", "desc": "按严重级别（🔴高 / 🟠中 / 🔵低）列出问题、风险与影响"},
            {"heading": "结论", "desc": "总结核心发现与工程判断"},
            {"heading": "建议与后续行动", "desc": "可执行建议、优先级与责任方"},
            {"heading": "附录", "desc": "术语表、参考文档、明细数据"},
        ],
    },
    "impact": {
        "label": "变更影响分析报告",
        "keywords": ["变更影响", "影响分析", "变更评估", "变更范围", "影响评估", "变更分析"],
        "sections": [
            {"heading": "报告信息", "desc": "文档编号、版本、日期、编制/审核、密级（用表格呈现）"},
            {"heading": "摘要", "desc": "一段话浓缩：变更内容、影响规模与核心结论"},
            {"heading": "变更描述", "desc": "变更源、变更内容、变更原因（变更前/后对比）"},
            {"heading": "分析范围与方法", "desc": "搜索深度、传播方向、关系类型、数据来源"},
            {"heading": "直接影响分析", "desc": "直接影响元素清单（表格：元素/类型/关系/影响度/级别）"},
            {"heading": "间接影响与传播路径", "desc": "沿关系链的分层传播、跨领域影响分析"},
            {"heading": "影响程度与风险分级", "desc": "影响统计、热力排序与风险（🔴高 / 🟠中 / 🔵低）"},
            {"heading": "处置建议", "desc": "逐风险应对动作、优先级与责任方"},
            {"heading": "结论", "desc": "变更可行性结论与关注事项"},
            {"heading": "附录", "desc": "完整影响明细、参考依据"},
        ],
    },
    "review": {
        "label": "预评审报告",
        "keywords": ["评审报告", "预评审", "审查", "合规", "规范性检查", "一致性检查"],
        "sections": [
            {"heading": "报告信息", "desc": "文档编号、版本、日期、编制/审核、密级（用表格呈现）"},
            {"heading": "摘要", "desc": "评审范围、总体结论与关键发现"},
            {"heading": "评审范围与依据", "desc": "评审对象、标准与依据"},
            {"heading": "规范符合性检查", "desc": "本体类型/关系约束与命名规范检查结果"},
            {"heading": "一致性问题清单", "desc": "问题描述、位置与类型"},
            {"heading": "严重级别汇总", "desc": "按严重级别统计与排序"},
            {"heading": "修复建议", "desc": "逐问题的可执行修复建议"},
            {"heading": "结论", "desc": "评审结论与放行建议"},
        ],
    },
}
DEFAULT_REPORT_TYPE = "analysis"


class ReportGenerator:
    """通用报告生成器（无状态，可被对话与工坊复用）。"""

    # 默认分节模板（LLM 按节聚焦生成，结构可控）
    DEFAULT_SECTIONS = [
        {"heading": "概述", "desc": "任务背景、范围与目标"},
        {"heading": "现状分析", "desc": "基于检索结果与上传资料的事实梳理"},
        {"heading": "详细分析", "desc": "核心内容：指标、对比、问题与影响"},
        {"heading": "结论", "desc": "总结核心发现"},
        {"heading": "建议", "desc": "下一步行动建议"},
    ]

    def detect_report_type(self, topic: str) -> str:
        """SP-R：按关键词识别报告类型（命中最多者优先）；无命中 → analysis。"""
        topic = topic or ""
        best_key, best_hits = DEFAULT_REPORT_TYPE, 0
        for key, cfg in REPORT_TYPES.items():
            hits = sum(1 for kw in cfg["keywords"] if kw in topic)
            if hits > best_hits:
                best_key, best_hits = key, hits
        return best_key

    def extract_custom_sections(self, topic: str) -> Optional[list]:
        """SP-R：解析用户自定义大纲——「按/以 … 章节/结构/顺序/形式 [:：] A/B/C」或「章节[:：] A/B/C」。

        分隔符支持 / 、/ 顿号、逗号；解析失败返回 None（回退类型模板）。
        """
        topic = (topic or "").strip()
        patterns = [
            # 支持「按 引言、问题分析、建议 三节生成」（无冒号）与「按以下章节：A/B/C」（有冒号）
            r"(?:按|以)(?:以下|如下)?(?:章节|结构|顺序|大纲|形式|模块)?\s*[为:：是]?\s*([^\n。；;]{2,80})",
            r"(?:章节|结构|大纲|形式)\s*[为:：是]\s*([^\n。；;]{2,80})",
            r"(?:自定义)?报告(?:章节|结构|形式|大纲)[:：]\s*([^\n。；;]{2,80})",
        ]
        for pat in patterns:
            m = re.search(pat, topic)
            if m:
                raw = m.group(1)
                # 清洗尾部「数量词+节+动词」（如「 三节生成」）与纯结尾动词；
                # 数量词前必须为空白，避免误删独立章节名（如「一、二、三」的「三」）
                raw = re.sub(r"\s+[一二三四五六七八九十两几多]\s*节?\s*(?:来)?(?:生成|撰写|编写|输出|制作|组织|写)?$", "", raw)
                raw = re.sub(r"\s*(?:来)?(?:生成|撰写|编写|输出|制作|组织|写)$", "", raw)
                parts = [p.strip() for p in re.split(r"[/、,，/；;|]", raw) if p.strip()]
                if len(parts) >= 2:
                    return [{"heading": p, "desc": ""} for p in parts]
        return None

    def sections_for(self, report_type: str, topic: str = "") -> list:
        """SP-R：自定义大纲优先 → 类型模板 → 默认模板。"""
        custom = self.extract_custom_sections(topic or "")
        if custom:
            return custom
        cfg = REPORT_TYPES.get(report_type)
        return cfg["sections"] if cfg else self.DEFAULT_SECTIONS

    def build_prompt(self, topic: str, context: str = "", sections: Optional[list] = None,
                     report_type: Optional[str] = None) -> str:
        """构造结构化报告生成提示（LLM 按节输出 markdown，写入规范对标行业工程报告）。"""
        secs = sections or self.DEFAULT_SECTIONS
        sec_block = "\n".join(f"### {s['heading']}\n（{s.get('desc','')}）" for s in secs)
        type_label = REPORT_TYPES.get(report_type or "", {}).get("label", "分析报告")
        custom_hint = "（章节结构按用户指定执行）" if self.extract_custom_sections(topic or "") else ""
        return (
            f"请基于以下素材撰写一份专业{type_label}《{topic}》{custom_hint}。\n"
            f"写作规范（对标行业工程报告标准）：\n"
            f"1. 语言正式严谨，面向工程评审与归档场景，避免口语化；摘要用一段话浓缩背景、方法与核心结论；\n"
            f"2. 涉及清单/对比/统计的数据一律用 markdown 表格呈现，指标需量化、表头简洁明确；\n"
            f"3. 风险与问题按严重程度分级（🔴高 / 🟠中 / 🔵低），逐条给出影响与建议；\n"
            f"4. 引用素材中的数据须与素材一致，不得虚构；素材不足以支撑处如实标注；\n"
            f"5. 结尾给出明确结论与可执行建议（含优先级）；「报告信息」节用表格给出文档编号/版本/日期/编制/审核。\n\n"
            f"严格按下列章节组织，每节用 markdown 三级标题（###），涉及数据用 markdown 表格"
            f"（前端可转图表）：\n\n"
            f"{sec_block}\n\n"
            f"素材：\n{context[:6000]}\n"
        )

    # 报告文档控制信息（封面/元信息条/导出用）
    _DOC_ABBR = {"analysis": "ENGR", "impact": "CIA", "review": "REV", "other": "RPT"}

    def build_meta(self, title: str = "", report_type: str = "") -> dict:
        """生成文档控制元信息：文档编号/版本/日期/状态/编制/密级。"""
        try:
            import zlib
            from datetime import datetime
            today = datetime.now()
            abbr = self._DOC_ABBR.get(report_type or "", "RPT")
            seq = zlib.crc32(str(title or "").encode("utf-8")) % 10000
            return {
                "doc_no": f"{abbr}-{today.strftime('%Y%m%d')}-{seq:04d}",
                "version": "V1.0",
                "date": today.strftime("%Y-%m-%d"),
                "status": "评审稿",
                "author": "AI 建模助手",
                "reviewer": "待审核",
                "classification": "内部",
            }
        except Exception:
            return {}        

    def structure(self, title: str, markdown: str, report_type: Optional[str] = None) -> dict:
        """把 LLM 生成的 markdown 报告解析为结构化 Report（sections + summary）。

        无章节标题（如 Mock 兜底输出）→ 全文作为单节「正文」，保证结构可用。
        """
        sections = []
        cur = None
        for line in (markdown or "").split("\n"):
            line = line.rstrip()
            if line.startswith("### "):
                if cur:
                    sections.append(cur)
                cur = {"heading": line[4:].strip(), "body": []}
            elif line.startswith("## "):
                if cur:
                    sections.append(cur)
                cur = {"heading": line[3:].strip(), "body": []}
            elif cur is not None:
                cur["body"].append(line)
        if cur:
            sections.append(cur)
        for s in sections:
            s["body"] = "\n".join(s["body"]).strip()[:3000]
        if not sections and (markdown or "").strip():
            sections = [{"heading": "正文", "body": (markdown or "").strip()[:3000]}]
        rtype = report_type or self.detect_report_type(title)
        return {
            "title": title,
            "sections": sections,
            "summary": (sections[-1]["body"] if sections else ""),
            "report_type": rtype,
            "meta": self.build_meta(title, rtype),
            "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        }

    def export_markdown(self, report: dict) -> str:
        """Report → 完整 markdown 文本（下载/存档用）。"""
        out = [f"# {report.get('title','')}", ""]
        meta = report.get("meta") or {}
        if meta:
            out += ["| 文档编号 | 版本 | 日期 | 密级 | 编制 |", "| --- | --- | --- | --- | --- |",
                    f"| {meta.get('doc_no','-')} | {meta.get('version','-')} | {meta.get('date','-')} | "
                    f"{meta.get('classification','-')} | {meta.get('author','-')} |", ""]
        for s in report.get("sections", []):
            out.append(f"## {s['heading']}")
            if s.get("body"):
                out.append(s["body"])
            if s.get("table"):
                out.append(s["table"])
            out.append("")
        return "\n".join(out)

    # ═══════════════ 多格式文档导出：Markdown / Word(.docx) / PDF ═══════════════
    # 单一数据契约 Report → 任意格式；依赖 python-docx / reportlab（已内置）。
    # body 支持轻量 markdown：**粗体**、表格（| a | b | 分隔行）、列表（- / 1.）、### 子标题。

    @staticmethod
    def _md_table_rows(table_md: str) -> list:
        """markdown 表格文本 → [[cell, ...], ...]（跳过分隔行 |---|---|）。"""
        rows = []
        for line in (table_md or "").splitlines():
            line = line.strip()
            if not line.startswith("|"):
                continue
            cells = [c.strip() for c in line.strip("|").split("|")]
            if cells and all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                continue  # 分隔行
            if cells:
                rows.append(cells)
        return rows

    @staticmethod
    def _split_blocks(markdown: str) -> list:
        """markdown 文本 → blocks：[{type:para|li|h3|table, ...}]。"""
        blocks, i, lines = [], 0, (markdown or "").splitlines()
        while i < len(lines):
            line = lines[i].rstrip()
            s = line.strip()
            if not s:
                i += 1
                continue
            if s.startswith("|"):
                rows, j = [], i
                while j < len(lines) and lines[j].strip().startswith("|"):
                    rows.append(lines[j].strip())
                    j += 1
                cells = ReportGenerator._md_table_rows("\n".join(rows))
                if cells:
                    blocks.append({"type": "table", "rows": cells})
                i = j
                continue
            if s.startswith("### ") or s.startswith("## "):
                blocks.append({"type": "h3", "text": s.lstrip("#").strip()})
            elif re.match(r"^[-*+]\s+", s) or re.match(r"^\d+[.、]\s+", s):
                blocks.append({"type": "li", "text": re.sub(r"^[-*+]\s+|^\d+[.、]\s+", "", s)})
            else:
                blocks.append({"type": "para", "text": s})
            i += 1
        return blocks

    # ── Word .docx 导出（python-docx：封面页 + 文档控制 + 页脚页码 + 样式化表格）──
    _DOCX_HDR_FILL = "1F4E79"   # 表头底色（品牌深蓝）
    _DOCX_ALT_FILL = "F2F6FB"   # 斑马纹

    @staticmethod
    def _docx_shade_cell(cell, color: str) -> None:
        """单元格底色（w:shd）。"""
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd")
        if shd is None:
            shd = cell._tc._element.makeelement(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd", {})
            tcPr.append(shd)
        shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "clear")
        shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", color)

    @staticmethod
    def _docx_footer_page(section) -> None:
        """页脚居中页码字段（PAGE）。"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fld1 = run._r.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldChar",
            {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType": "begin"})
        instr = run._r.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instrText",
            {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space": "preserve"})
        instr.text = "PAGE"
        fld2 = run._r.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldChar",
            {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType": "end"})
        run._r.append(fld1)
        run._r.append(instr)
        run._r.append(fld2)
        from docx.shared import Pt, RGBColor
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    def _docx_heading(self, doc, text: str, level: int = 1) -> None:
        """带品牌色的标题（标题字体 eastAsia 同步）。"""
        from docx.shared import Pt, RGBColor
        from docx.oxml.ns import qn
        h = doc.add_heading(text or "", level)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            run.font.size = Pt({1: 16, 2: 14, 3: 12}.get(level, 11))

    def _docx_table(self, doc, rows: list) -> None:
        """样式化表格：表头深蓝白字 + 斑马纹 + 边框。"""
        from docx.shared import Pt, RGBColor
        from docx.oxml.ns import qn
        if not rows:
            return
        cols = len(rows[0])
        tb = doc.add_table(rows=len(rows), cols=cols)
        try:
            tb.style = "Table Grid"
        except Exception:
            pass
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row[:cols]):
                c = tb.rows[ri].cells[ci]
                c.text = ""
                p = c.paragraphs[0]
                run = p.add_run(str(cell))
                run.font.size = Pt(9)
                run.font.name = "宋体"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                if ri == 0:
                    self._docx_shade_cell(c, self._DOCX_HDR_FILL)
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                elif ri % 2 == 0:
                    self._docx_shade_cell(c, self._DOCX_ALT_FILL)

    def export_docx(self, report: dict, path: str | None = None):
        """Report → Word 文档（封面页 + 文档控制表 + 样式化正文）。path 为空返回 bytes。"""
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        doc = Document()
        # 默认正文字体（中文）
        st = doc.styles["Normal"]
        st.font.name = "宋体"
        st.font.size = Pt(10.5)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        self._docx_footer_page(doc.sections[0])

        title = (report.get("title") or "报告").strip() or "报告"
        meta = report.get("meta") or {}
        type_label = REPORT_TYPES.get(report.get("report_type") or "", {}).get("label", "工程报告")

        # ── 封面页 ──
        for _ in range(4):
            doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.size = Pt(22)
        run.bold = True
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(f"{type_label}（AI 建模生成）")
        run2.font.size = Pt(13)
        run2.font.color.rgb = RGBColor(0x7F, 0x8C, 0x9E)
        run2.font.name = "微软雅黑"
        run2._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        for _ in range(6):
            doc.add_paragraph()
        if meta:
            meta_rows = [("文档编号", meta.get("doc_no", "-")), ("版本", meta.get("version", "-")),
                         ("日期", meta.get("date", "-")), ("密级", meta.get("classification", "-")),
                         ("编制", meta.get("author", "-")), ("审核", meta.get("reviewer", "-")),
                         ("状态", meta.get("status", "-"))]
            tb = doc.add_table(rows=len(meta_rows), cols=2)
            try:
                tb.style = "Table Grid"
            except Exception:
                pass
            for i, (k, v) in enumerate(meta_rows):
                ck, cv = tb.rows[i].cells[0], tb.rows[i].cells[1]
                ck.text = ""
                rk = ck.paragraphs[0].add_run(k)
                rk.font.name = "微软雅黑"
                rk._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                rk.font.size = Pt(10.5)
                rk.bold = True
                self._docx_shade_cell(ck, "E8EEF6")
                cv.text = v
                rv = cv.paragraphs[0].runs[0]
                rv.font.size = Pt(10.5)
                rv.font.name = "微软雅黑"
                rv._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        doc.add_page_break()

        # ── 正文分节 ──
        for s in report.get("sections", []):
            self._docx_heading(doc, s.get("heading") or "", 1)
            for blk in self._split_blocks(s.get("body") or ""):
                if blk["type"] == "table":
                    self._docx_table(doc, blk["rows"])
                elif blk["type"] == "h3":
                    self._docx_heading(doc, blk["text"], 3)
                elif blk["type"] == "li":
                    self._docx_rich(doc.add_paragraph(style="List Bullet"), blk["text"])
                else:
                    self._docx_rich(doc.add_paragraph(), blk["text"])
            if s.get("table"):
                rows = self._md_table_rows(s["table"])
                if rows:
                    self._docx_table(doc, rows)
        return self._docx_bytes(doc, path)

    @staticmethod
    def _docx_rich(paragraph, text: str) -> None:
        """段落富文本：**粗体** 拆分为加粗 run。"""
        pos = 0
        for m in re.finditer(r"\*\*(.+?)\*\*", text or ""):
            if m.start() > pos:
                paragraph.add_run(text[pos:m.start()])
            paragraph.add_run(m.group(1)).bold = True
            pos = m.end()
        if pos < len(text or ""):
            paragraph.add_run(text[pos:])

    @staticmethod
    def _docx_bytes(doc, path: str | None):
        import io
        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()
        if path:
            os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
            with open(path, "wb") as f:
                f.write(data)
            return path
        return data

    # ── PDF 导出（reportlab platypus，自动注册系统中文字体）──
    _CN_FONT_CACHE: dict = {}

    @classmethod
    def _cn_font(cls) -> Optional[str]:
        """注册系统中文字体（黑体/宋体/雅黑），返回字体名；无则 None（兜底默认字体）。"""
        if cls._CN_FONT_CACHE:
            return cls._CN_FONT_CACHE.get("name")
        name = None
        try:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            cands = [
                (r"C:\Windows\Fonts\simhei.ttf", "SimHei"),
                (r"C:\Windows\Fonts\simsun.ttc", "SimSun"),
                (r"C:\Windows\Fonts\msyh.ttc", "MSYH"),
            ]
            for fpath, fname in cands:
                if os.path.exists(fpath):
                    try:
                        pdfmetrics.registerFont(TTFont(fname, fpath))
                        name = fname
                        break
                    except Exception:
                        continue
        except Exception:
            name = None
        cls._CN_FONT_CACHE["name"] = name
        return name

    @staticmethod
    def _pdf_rich(text: str, font: str | None) -> str:
        """pdf 富文本：**粗体** → <b>，其余 XML 转义，\n → <br/>。"""
        def _esc(t: str) -> str:
            return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        parts, pos = [], 0
        for m in re.finditer(r"\*\*(.+?)\*\*", text or ""):
            parts.append(_esc(text[pos:m.start()]))
            parts.append(f"<b>{_esc(m.group(1))}</b>")
            pos = m.end()
        parts.append(_esc(text[pos:]))
        out = "".join(parts)
        if font:
            return out.replace("\n", "<br/>")
        return out.replace("\n", "<br/>")

    # ── PDF 导出（reportlab platypus：封面页 + 页脚页码 + 样式化表格，自动注册系统中文字体）──
    _PDF_HDR_FILL = "1F4E79"
    _PDF_ALT_FILL = "F2F6FB"

    def export_pdf(self, report: dict, path: str | None = None):
        """Report → PDF（封面页 + 分节正文 + 页脚页码）。path 为空返回 bytes。"""
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                        TableStyle, PageBreak)
        from reportlab.lib.enums import TA_CENTER, TA_LEFT

        font = self._cn_font()
        base = getSampleStyleSheet()
        fname = font or base["BodyText"].fontName
        st_cover = ParagraphStyle("CnCover", parent=base["Title"], fontName=font or base["Title"].fontName,
                                  fontSize=20, leading=28, alignment=TA_CENTER, textColor=colors.HexColor("#1F4E79"))
        st_cover_sub = ParagraphStyle("CnCoverSub", parent=base["BodyText"], fontName=fname,
                                      fontSize=12, leading=18, alignment=TA_CENTER, textColor=colors.HexColor("#7F8C9E"))
        st_h1 = ParagraphStyle("CnH1", parent=base["Heading1"], fontName=font or base["Heading1"].fontName,
                               fontSize=14, leading=18, textColor=colors.HexColor("#1F4E79"))
        st_h3 = ParagraphStyle("CnH3", parent=base["Heading3"], fontName=font or base["Heading3"].fontName,
                               fontSize=11, leading=15, textColor=colors.HexColor("#1F4E79"))
        st_body = ParagraphStyle("CnBody", parent=base["BodyText"], fontName=fname,
                                 fontSize=10, leading=15, alignment=TA_LEFT)
        st_th = ParagraphStyle("CnTh", parent=st_body, fontSize=9, leading=12, textColor=colors.white)
        st_td = ParagraphStyle("CnTd", parent=st_body, fontSize=9, leading=12)

        def _pdf_table(rows: list) -> Table:
            data = []
            for ri, row in enumerate(rows):
                cells = [Paragraph(self._pdf_rich(c, font), st_th if ri == 0 else st_td) for c in row]
                data.append(cells)
            tb = Table(data, repeatRows=1, hAlign="LEFT")
            style = [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B9C4D8")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(self._PDF_HDR_FILL)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
            for i in range(2, len(rows), 2):
                style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(self._PDF_ALT_FILL)))
            tb.setStyle(TableStyle(style))
            return tb

        import io
        buf = io.BytesIO()
        title = (report.get("title") or "报告").strip() or "报告"
        meta = report.get("meta") or {}
        type_label = REPORT_TYPES.get(report.get("report_type") or "", {}).get("label", "工程报告")

        def _footer(canvas, doc):
            canvas.saveState()
            if doc.page > 1:  # 封面不显示页码
                canvas.setFont(font or "Helvetica", 8)
                canvas.setFillColor(colors.HexColor("#888888"))
                canvas.drawCentredString(A4[0] / 2, 9 * mm, f"- {doc.page - 1} -")
            canvas.restoreState()

        doc = SimpleDocTemplate(buf, pagesize=A4,
                                leftMargin=18 * mm, rightMargin=18 * mm,
                                topMargin=16 * mm, bottomMargin=16 * mm,
                                title=title)
        story = []
        # ── 封面页 ──
        story += [Spacer(1, 70 * mm), Paragraph(title, st_cover),
                  Spacer(1, 10), Paragraph(f"{type_label}（AI 建模生成）", st_cover_sub), Spacer(1, 30 * mm)]
        if meta:
            meta_rows = [("文档编号", meta.get("doc_no", "-")), ("版本", meta.get("version", "-")),
                         ("日期", meta.get("date", "-")), ("密级", meta.get("classification", "-")),
                         ("编制", meta.get("author", "-")), ("审核", meta.get("reviewer", "-")),
                         ("状态", meta.get("status", "-"))]
            data = [[Paragraph(k, st_td), Paragraph(v, st_td)] for k, v in meta_rows]
            mtb = Table(data, colWidths=[36 * mm, None], hAlign="CENTER")
            mtb.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B9C4D8")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EEF6")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.append(mtb)
        story.append(PageBreak())
        # ── 正文分节 ──
        for s in report.get("sections", []):
            story.append(Paragraph(s.get("heading") or "", st_h1))
            story.append(Spacer(1, 6))
            for blk in self._split_blocks(s.get("body") or ""):
                if blk["type"] == "table":
                    story.append(_pdf_table(blk["rows"]))
                elif blk["type"] == "h3":
                    story.append(Paragraph(blk["text"], st_h3))
                elif blk["type"] == "li":
                    story.append(Paragraph("• " + self._pdf_rich(blk["text"], font), st_body))
                else:
                    story.append(Paragraph(self._pdf_rich(blk["text"], font), st_body))
                story.append(Spacer(1, 3))
            if s.get("table"):
                rows = self._md_table_rows(s["table"])
                if rows:
                    story.append(_pdf_table(rows))
            story.append(Spacer(1, 8))
        doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
        data = buf.getvalue()
        if path:
            os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
            with open(path, "wb") as f:
                f.write(data)
            return path
        return data

    def export(self, report: dict, fmt: str = "md", path: str | None = None):
        """统一导出出口：fmt ∈ md | docx | pdf。path 为空返回 bytes，否则写文件返回路径。"""
        fmt = str(fmt or "md").lower().lstrip(".")
        if fmt == "docx":
            return self.export_docx(report, path)
        if fmt == "pdf":
            return self.export_pdf(report, path)
        if fmt != "md":
            raise ValueError(f"不支持的导出格式: {fmt}")
        data = self.export_markdown(report).encode("utf-8")
        if path:
            os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
            with open(path, "wb") as f:
                f.write(data)
            return path
        return data


# 模块级单例（与 llm_client 同风格）
report_generator = ReportGenerator()
