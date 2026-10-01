# -*- coding: utf-8 -*-
"""AgentPipeline Mixin：附件加载、历史检索、话题标注与上下文预算。

由 tools/split_pipeline.py 从 agent/pipeline.py 机械切分而成；⚠️ 切分脚本**已一次性执行完毕、不可重跑**—— 此后本文件按普通源码维护（方法体与其它模块一样可直接改）。"""
from .common import *

# P2 视觉通道（2026-09-21）：图片附件扩展名（与 routers/conversations.py 的 IMAGE_EXTS 保持同一集合）
_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}


def ctx_budget_tokens(kind: str, fallback_char_key: str, fallback: int) -> int:
    """P1-4b（2026-09-21）上下文预算取值：**占比制优先 → 绝对值 → 字符版兜底**。

    kind ∈ {'retrieval','history'}。占比键 `context.budget_{kind}_ratio` > 0 时，
    以「budget_window_tokens × 占比」为准（换模型自适应，依据调研 §4.3）；
    否则回落 T6 绝对值键 `budget_{kind}_tokens`（0 则再回落字符版键）。
    与既有「token 项填 0 回退字符版」的兼容模式同构。任何异常返回 fallback。
    """
    try:
        from core import config as _cfg
        ratio = float(_cfg.get("context", f"budget_{kind}_ratio", 0) or 0)
        if ratio > 0:
            win = int(_cfg.get("context", "budget_window_tokens", 65536) or 65536)
            return max(16, int(win * ratio))
        return (int(_cfg.get("context", f"budget_{kind}_tokens", 0))
                or int(_cfg.get("context", fallback_char_key, fallback)))
    except Exception:
        return fallback


class HistoryMixin:
    """附件加载、历史检索、话题标注与上下文预算。"""

    def _load_attachment_text(self, attachments):
        """返回 (parsed_blocks, parsed_count, skipped_names)。parsed_blocks 为 ["[文件名]\n内容…"] 列表。"""
        blocks, parsed, skipped = [], 0, []
        for a in attachments or []:
            url = (a.get("url") or "") if isinstance(a, dict) else ""
            fname = (a.get("filename") or url.split("/")[-1] or "附件") if isinstance(a, dict) else "附件"
            if not url.startswith("/static/"):
                continue
            rel = url.replace("/static/", "")
            path = os.path.join(STATIC_DIR, rel)
            if not os.path.exists(path):
                skipped.append(fname)
                continue
            ext = os.path.splitext(fname)[1].lower()
            text = ""
            try:
                if ext in (".txt", ".md", ".csv", ".json", ".xml", ".log", ".sysml", ".yaml", ".yml"):
                    raw = open(path, "rb").read()
                    for enc in ("utf-8", "gbk", "latin-1"):
                        try:
                            text = raw.decode(enc)
                            break
                        except Exception:
                            continue
                elif ext == ".pdf":
                    import pdfplumber
                    with pdfplumber.open(path) as pdf:
                        text = "\n".join((p.extract_text() or "") for p in pdf.pages[:20])
                elif ext == ".docx":
                    import docx
                    d = docx.Document(path)
                    text = "\n".join(p.text for p in d.paragraphs[:300])
                elif ext == ".xlsx":
                    import openpyxl
                    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                    rows = []
                    for ws in wb.worksheets[:3]:
                        for row in ws.iter_rows(max_row=50, values_only=True):
                            rows.append(" | ".join(str(c) for c in row if c is not None))
                    text = "\n".join(rows)
            except Exception:
                skipped.append(fname)
                continue
            text = (text or "").strip()
            if not text:
                skipped.append(fname)
                continue
            # 全文按段分块（每块 ≤2500 字符），保证长文档后半部分同样参与检索与上下文注入
            for i, seg in enumerate(text[i:i + 2500] for i in range(0, len(text), 2500)):
                # 2026-09-17 S3：解析上限维持 40（曾试降到 8，但该值决定附件语义召回候选池 retrieval_att，
                #                 降它只省解析耗时、不省 token，故回退；token 侧改由 _ATT_INJECT_CAP 在注入点封顶）
                if len(blocks) >= 40:
                    break
                head = f"【{fname}】" if i == 0 else f"【{fname}#{i + 1}】"
                blocks.append(f"{head}\n{seg}")
            parsed += 1
        return blocks, parsed, skipped

    # ── P2 视觉通道（2026-09-21）：图片附件 → 多模态 content 块 ───────────────
    #   背景（实测）：`_load_attachment_text` 的后缀白名单**不含任何图片格式**，
    #   图片 → text="" → 落进 skipped → att_blocks=[] → prompt 里连文件名都没有
    #   → **模型从未见过图**（会话里传架构图，AI 收到的是一片空白）。
    #   本节把图片按 OpenAI 多模态 image_url 块注入 user 消息，让图真正进入模型。
    def _load_attachment_images(self, attachments, max_images=2, max_side=1280,
                                max_bytes=4194304, jpeg_quality=85):
        """把图片附件读为 OpenAI 多模态 image_url 块（base64 data URL）。

        返回 (blocks, loaded, skipped)：
          blocks  —— [{"type":"image_url","image_url":{"url":"data:image/png;base64,..."}}]
          loaded  —— 成功载入的文件名（用于在 prompt 里告知模型「有图可看」）
          skipped —— ["文件名（原因）"]，**必须留痕**（不得静默丢弃，这是本通道的纪律）

        为什么要重编码而不是直接 base64 原图：
          · 视觉 token 与**像素数**成正比、与压缩率无关 → 直接传原图（如 4000px 截图）
            会白烧几千 token，且大图对架构图这类线框内容的识别收益接近于零；
          · 统一压到长边 max_side 后再编码，token 可控、体积可控（20MB 上限的附件
            直接内联进请求体会把请求体撑爆）。
        逐图防护：路径必须落在 STATIC_DIR 下且存在 → PIL 能打开 → 重编码后 ≤ max_bytes。
        任一不满足**只跳过该图**并给出原因，不影响其他图与主链路。
        """
        blocks, loaded, skipped = [], [], []
        try:
            import io as _io
            import base64 as _b64
            from PIL import Image
        except Exception as e:
            return blocks, loaded, ["（图片能力不可用：缺少 Pillow）"] if attachments else []
        cap = max(0, int(max_images or 0))
        for a in attachments or []:
            if len(blocks) >= cap:
                break
            url = (a.get("url") or "") if isinstance(a, dict) else ""
            fname = (a.get("filename") or url.split("/")[-1] or "附件") if isinstance(a, dict) else "附件"
            if os.path.splitext(fname)[1].lower() not in _IMAGE_EXTS:
                continue  # 非图片附件归文本通道处理，不算「被跳过」
            if not url.startswith("/static/"):
                skipped.append(f"{fname}（附件地址非 /static/ 前缀，无法定位落盘副本）")
                continue
            path = os.path.join(STATIC_DIR, url.replace("/static/", ""))
            if not os.path.exists(path):
                skipped.append(f"{fname}（落盘副本不存在）")
                continue
            try:
                with Image.open(path) as im:
                    im.load()
                    w0, h0 = im.size
                    has_alpha = im.mode in ("RGBA", "LA") or (
                        im.mode == "P" and "transparency" in im.info)
                    if max(w0, h0) > max_side:  # 等比缩小（视觉 token ∝ 像素数）
                        r = float(max_side) / float(max(w0, h0))
                        im = im.resize((max(1, int(w0 * r)), max(1, int(h0 * r))))
                    buf = _io.BytesIO()
                    if has_alpha:  # 有透明通道一律 PNG，避免 JPEG 把透明压成黑底
                        im.convert("RGBA").save(buf, format="PNG", optimize=True)
                        mime = "image/png"
                    else:
                        im.convert("RGB").save(buf, format="JPEG",
                                               quality=int(jpeg_quality or 85), optimize=True)
                        mime = "image/jpeg"
                    raw = buf.getvalue()
            except Exception as e:
                skipped.append(f"{fname}（图片解码/编码失败：{type(e).__name__}）")
                continue
            if len(raw) > max_bytes:
                skipped.append(f"{fname}（重编码后 {len(raw) // 1024}KB 超过 {max_bytes // 1024}KB 上限）")
                continue
            blocks.append({"type": "image_url",
                           "image_url": {"url": f"data:{mime};base64," + _b64.b64encode(raw).decode("ascii")}})
            loaded.append(fname)
        return blocks, loaded, skipped

    def _prepare_attachment_vision(self, attachments, provider_id=None, conn=None):
        """视觉通道统一入口：决定「这次要不要把图片喂给模型」，并给出**可审计留痕**。

        返回 (blocks, info)：
          blocks —— 可直接并入 user 消息的多模态块（不需要时为空列表）
          info   —— 留痕字典（并入 attachments_info.vision，随响应可见）：
                    enabled / provider_vision / reason / loaded / skipped / images

        三重条件**全满足**才注入（任一不满足 → 不注入，但逐条写明原因）：
          ① `vision.enabled` = true（总开关，默认 false）
          ② 所选 provider 标了视觉能力（`llm.provider_supports_vision`）
          ③ 附件里确实有图片文件

        纪律：条件不满足时**绝不静默**。老实现是把图片塞进 skipped 就完事（连"为什么"都没有），
        本通道必须把原因写进 info 落进响应——否则「图没生效」与「模型看图了但没看懂」
        在现象上无法区分（同族教训：静默兜底会把真故障藏起来）。
        """
        info = {"enabled": False, "provider_vision": False, "reason": "",
                "loaded": [], "skipped": [], "images": 0}
        try:
            imgs = [a for a in (attachments or [])
                    if isinstance(a, dict)
                    and os.path.splitext(a.get("filename") or a.get("url") or "")[1].lower() in _IMAGE_EXTS]
            if not imgs:
                return [], info
            info["images"] = len(imgs)
            from core import config as _cfg
            enabled = bool(_cfg.as_bool("vision", "enabled", False))
            info["enabled"] = enabled
            from llm import resolve_provider_cfg, provider_supports_vision
            cfg = resolve_provider_cfg(provider_id, conn=conn)
            vision_ok = provider_supports_vision(cfg)
            info["provider_vision"] = vision_ok
            if not enabled:
                info["reason"] = "视觉通道未启用（配置 vision.enabled=false）"
                info["skipped"] = [f"{a.get('filename') or '图片'}（{info['reason']}）" for a in imgs]
                return [], info
            if not vision_ok:
                _who = cfg.get("name") or (f"provider#{cfg.get('id')}" if cfg.get("id") else "默认模型")
                info["reason"] = (f"当前模型未声明图像输入能力（{_who}；"
                                  "需 model_type=vision 或 tags 含 vision）")
                info["skipped"] = [f"{a.get('filename') or '图片'}（{info['reason']}）" for a in imgs]
                return [], info
            blocks, loaded, skipped = self._load_attachment_images(
                imgs,
                max_images=int(_cfg.get("vision", "max_images", 2) or 2),
                max_side=int(_cfg.get("vision", "max_side", 1280) or 1280),
                max_bytes=int(_cfg.get("vision", "max_bytes", 4194304) or 4194304),
                jpeg_quality=int(_cfg.get("vision", "jpeg_quality", 85) or 85))
            info["loaded"] = loaded
            info["skipped"] = skipped
            info["reason"] = f"已注入 {len(blocks)} 张" if blocks else "图片均未能载入（见 skipped）"
            return blocks, info
        except Exception as e:
            info["reason"] = f"视觉通道异常（已降级为不注入）：{type(e).__name__}: {e}"
            try:
                logger.warning("[vision] %s", info["reason"])
            except Exception:
                pass
            return [], info

    # ── 会话历史注入 v2：话题感知组装（P0+P1，替代纯时间窗口）──
    #   ① 当前话题原文段（最近 N 条）+ 上一话题切换边界（2 条）→ 原文逐字
    #   ② 语义拉回：与当前输入相关、未被原文注入的历史片段（top-k 原文，解决切回旧话题）
    #   ③ 其他话题：分话题独立摘要（scope=topic:{label}，消除混合摘要稀释，缓存复用）
    #   预算驱动：①占 50% → ②到 75% → ③剩余（不足则低优先级停，不送全量）
    #   话题打标失败/无话题 → 回退 _load_history_legacy（三层窗口）
    #   P1-8（2026-10-01）「输入/输出都按相关性评估」三处收口：
    #     · ① 由「吃满为止」改**均分份额**（防单条长输出独占，用户输入不再被挤成残句）；
    #     · `used_ids` 只记**真正注入**的（旧写法的「被考虑过」会让消息彻底消失）；
    #     · ② 分角色配额（user/assistant 各有名额）+ 块头带轮次/角色 + 丢弃与当前输入逐字重复的块。
    def _load_history(self, conversation_id, limit=None, cur_input=None):
        from core import config as _cfg
        msg_max = int(_cfg.get("context", "history_msg_max_chars", 1500))
        summary_chars = int(_cfg.get("context", "history_summary_chars", 800))
        cur_max = int(_cfg.get("context", "topic_current_max_msgs", 8))
        boundary = int(_cfg.get("context", "topic_boundary_keep", 2))
        pull_topk = int(_cfg.get("context", "topic_retrieve_topk", 6))
        pull_th = float(_cfg.get("context", "topic_retrieve_threshold", 0.18))
        budget = int(_cfg.get("context", "budget_history_chars", 3000))
        # P1-8（2026-10-01）历史（输入/输出）相关性召回收口 —— 三个新旋钮
        pull_per_role = int(_cfg.get("context", "topic_retrieve_per_role_cap", 3))
        raw_min_share = int(_cfg.get("context", "history_raw_min_share_tokens", 96))
        drop_echo = _cfg.as_bool("context", "topic_retrieve_drop_echo", True)
        # P0-6（2026-09-30）：会话产物摘要块。多轮上下文此前只有「消息」一种载体 ——
        # sysml_versions / artifacts 里的产物**没有任何注入通道**（全仓读它们的生产代码
        # 只有 agent/utils.py，且全在归档写入路径）。实测会话 514 第 2 轮因此把第 1 轮
        # 建好的需求模型当成不存在。无产物 → _digest 为空列表 → 返回值与改动前逐字一致。
        _digest = []
        try:
            conn = get_db()
            rows = conn.execute(
                "SELECT id, role, content FROM messages WHERE conversation_id=? AND role IN ('user','assistant') "
                "ORDER BY id", (conversation_id,)).fetchall()
            topics = self._tag_topics(conn, conversation_id)
            try:
                from agent.session_artifacts import build_digest
                _dgt = build_digest(conn, conversation_id)
                if _dgt:
                    _digest = [{"role": "system", "content": _dgt}]
            except Exception:
                _digest = []
            conn.close()
        except Exception:
            return []
        if not rows:
            return _digest
        msgs = [{"id": r["id"], "role": r["role"], "content": (r["content"] or "")[:msg_max],
                 "topic": topics.get(r["id"], "")} for r in rows]
        if not any(m["topic"] for m in msgs):  # 打标失败 → 回退旧三层窗口
            return _digest + self._load_history_legacy(conversation_id, limit)
        # 按话题顺序分组（保持原顺序）
        groups = []
        for m in msgs:
            t = m["topic"]
            if groups and groups[-1]["topic"] == t:
                groups[-1]["msgs"].append(m)
            else:
                groups.append({"topic": t, "msgs": [m]})
        # 当前话题判定：当前输入未落库（execute 先回复后落库），不能用最后一组；
        # 用当前输入与各组代表向量相似度匹配（真 embedding 优先，dense 阈值独立量纲；
        # bigram 降级：≥0.12 取最佳组 → 切回旧话题时命中旧组；否则视为新话题取最后一组兜底）
        from collections import Counter
        from knowledge_engine import VectorEngine
        _ve = VectorEngine()
        cur_group = groups[-1]
        if cur_input:
            best_g, best_s = None, 0.0
            dense_th = float(_cfg.get("context", "topic_group_match_dense", 0.30))
            # dense：每组代表 = 段首 3 条消息（各截 200 字）拼接，与当前输入同批向量化
            reprs = [" ".join((m["content"] or "")[:200] for m in g["msgs"][:3]) for g in groups]
            scores, backend = self._semantic_scores(reprs, cur_input)
            if backend == "dense":
                for g, s in zip(groups, scores):
                    if s > best_s:
                        best_s, best_g = s, g
                if best_g is not None and best_s >= dense_th:
                    cur_group = best_g
            else:  # bigram 降级：段内全部消息 bigram 累加为段代表向量
                qv = _ve._vector(cur_input)
                for g in groups:
                    tv = Counter()
                    for m in g["msgs"]:
                        tv.update(_ve._vector(m["content"]))
                    s = _ve._cosine(qv, tv) if tv else 0.0
                    if s > best_s:
                        best_s, best_g = s, g
                if best_g is not None and best_s >= 0.12:
                    cur_group = best_g
        raw_list = []
        pull_blocks = []
        sum_blocks = []
        used = 0
        # T6：预算改 token 驱动（budget_history_tokens 优先，旧字符配置兜底）
        # P1-4b：改为占比制优先（context.budget_history_ratio > 0 时按窗口比例算）
        from core.token_counter import count_tokens as _ct
        budget_tok = ctx_budget_tokens("history", "budget_history_chars", 3000)
        raw_cap = int(budget_tok * 0.5)
        pull_cap = int(budget_tok * 0.75)

        def _clip(text, cap_tokens):
            """按 token 上限裁剪文本（字符比例逼近，中英混排近似）。"""
            if not text or _ct(text) <= cap_tokens:
                return text
            ratio = max(len(text) / max(_ct(text), 1), 0.5)
            cut = text[:int(max(cap_tokens, 16) * ratio)]
            for _ in range(2):
                if _ct(cut) <= cap_tokens:
                    return cut
                cut = cut[:int(len(cut) * cap_tokens / max(_ct(cut), 1))]
            return cut

        def push_raw(m, share):
            nonlocal used
            if not m["content"]:
                return
            room = raw_cap - used
            if room <= 0:
                return
            # 每条**最多** share：这才是「装几条」与「装多长」的分离点（见下方均分说明）
            c = _clip(m["content"], min(share, room))
            if not c:
                return
            raw_list.append({"role": m["role"], "content": c, "id": m["id"]})
            used += _ct(c)

        # ① 当前话题原文（最近 cur_max 条）：**均分份额**填充，防被单条长输出吃满。
        #   P1-8 实测（真库 conv=514，raw_cap=1000 tok）：旧实现「newest-first 吃满为止」⇒
        #   最新一条 4843 字符的助手回复独占 994 tok，紧接着的**最新用户输入**只拿到
        #   `max(1000-994,16)=16` tok（残句「按上述方案落地：在既有」），更早 3 条用户输入
        #   整条消失 —— 而用户输入才是「意图与约束」的载体（它们是短文本，本来就几乎不占预算）。
        #   均分后 8 条各得 raw_cap//n（≥ min_share）⇒ 近 N 轮连续可见；
        #   用不完的余量自然流向 ② 相关性通道（按相关度挑出来的内容更该吃预算）。
        _tail = list(cur_group["msgs"][-cur_max:])
        _bnd = []
        # 切换边界：当前组很短（正在切换话题）时补其前一组最后 boundary 条原文，保证语义连续
        if len(cur_group["msgs"]) <= 2 and cur_group is not groups[0]:
            gi = groups.index(cur_group)
            _bnd = list(groups[gi - 1]["msgs"][-boundary:])
        # P1-18：keep-last-N 轮（跨话题）verbatim —— 话题快速切换时，最近 N 轮不因话题边界被压成摘要。
        #   默认 0 = 关闭（零行为漂移）；>0 时全局最近 N 轮（从倒数第 N 个 user 消息起）无条件进原文候选。
        _keep_n = int(_cfg.get("context", "keep_last_n_turns", 0))
        _keep_msgs = []
        if _keep_n > 0:
            _uidx = [i for i, m in enumerate(msgs) if m["role"] == "user"]
            if _uidx:
                _start = _uidx[-_keep_n] if len(_uidx) >= _keep_n else _uidx[0]
                _keep_msgs = [m for m in msgs[_start:]]
        # 合并候选按 id 升序去重（时间序）；keep_n=0 时 _keep_msgs 空、_bnd/_tail 无重叠 → 严格等于旧 _bnd+_tail
        _cand = {}
        for m in (_keep_msgs + _bnd + _tail):
            _cand[m["id"]] = m
        _rcands = [_cand[k] for k in sorted(_cand)]
        # 两级分配：① 先按「人均份额」保**条数**（谁也不许独占）；② 余量再回填给被截断的
        #   （见下方第二轮）。只做 ① 会把短会话里唯一那条长回复白切一半；只做「吃满为止」
        #   就是本批要修的病灶。两者组合才同时满足「条数覆盖」与「不浪费」。
        _share = max(raw_min_share, raw_cap // max(len(_rcands), 1))
        for m in reversed(_rcands):                 # newest-first：预算真不够时优先保最新
            push_raw(m, _share)
        raw_list.reverse()  # 还原时间顺序（最近的在最后）
        # 第二轮：把**没用完**的额度回填给被截断的条目（newest-first，上不封顶至原文全长）。
        # 理由：① 的额度就是 50%（`raw_cap`），不用满等于把预算白扔；② 的额度是**另外**的 25%
        #   （`pull_cap = 75%`），① 用满后 ② 仍有 500 tok（约 3 条 300 字符片段），不是被 ① 抢走。
        # 实测两种取法的差别（真库 8 会话）：带回填 ① 用满 1000 tok、user 覆盖 36.2%；
        #   不回填 ① 只用 752 tok、user 覆盖 40.4%（② 多出 1 块）—— 差 1 块换「最新一条原文
        #   从 125 tok 涨到 ~370 tok」，取带回填（额度语义与设计一致，且不浪费）。
        if used < raw_cap and raw_list:
            _rmap = {m["id"]: m for m in _rcands}
            for _i in range(len(raw_list) - 1, -1, -1):   # raw_list 已时间序 ⇒ 倒序=newest-first
                _room = raw_cap - used
                if _room < raw_min_share:
                    break
                _src = _rmap.get(raw_list[_i].get("id"))
                if _src is None:
                    continue
                _cur_tok = _ct(raw_list[_i]["content"])
                _full = _src["content"] or ""
                if _ct(_full) <= _cur_tok:
                    continue                              # 本来就没被截断
                _c2 = _clip(_full, min(_ct(_full), _cur_tok + _room))
                _new_tok = _ct(_c2)
                if _new_tok > _cur_tok:
                    raw_list[_i]["content"] = _c2
                    used += _new_tok - _cur_tok
        used_ids = {m["id"] for m in raw_list if m.get("id")}
        # ⚠️ 此处**不得**再写 `used_ids.update(cur_group["msgs"][-cur_max:])`。
        #   那是「被考虑过」而非「已注入」。旧写法把因预算不足没进去的消息也标成已用
        #   ⇒ 它们既不在①原文、也不在②相关性拉回 ⇒ **彻底消失**。铁证（真库）：
        #   conv=1 里相似度 **0.988** 的用户输入、conv=325 里 **0.927** 的那条，都是这么丢的
        #   —— 被吞掉的恰恰是最该被拉回的内容。
        # ② 语义拉回（query=当前输入+当前话题，排除已注入原文；同话题域加权）
        query = f"{cur_input or ''} {cur_group['topic']}"
        rest = [m for m in msgs if m["id"] not in used_ids]
        # 逐字重复当前输入的历史消息 → **不参与相关性评估**（零信息量：模型当前输入里已有）。
        # 实测（真库 conv=1）：两条 13 字符的追问以 **0.988** 的高分占据拉回槽位，
        # 内容与当前输入一字不差 —— 高相似度 ≠ 有信息量。
        # ⚠️ 必须在**候选层**过滤，不能只丢块：它分数最高，先占掉一个配额槽再被丢弃，
        #    等于白瞎一个名额（把过滤放在这里，配额才算给了真有内容的片段）。
        _echo = re.sub(r"\s+", "", cur_input or "")
        if drop_echo and _echo:
            rest = [m for m in rest if re.sub(r"\s+", "", m.get("content") or "") != _echo]
        pulled = self._search_history(rest, query, pull_topk, pull_th,
                                      cur_topic=cur_group["topic"], per_role_cap=pull_per_role)
        # 轮次表（user 消息序号 = 轮次）：拉回块必须能被回指到「第几轮的谁说的」，
        # 而不是一句无出处的「相关历史片段」（延续 §7「可追溯」纪律）。
        _turn_of, _t = {}, 0
        for m in msgs:
            if m["role"] == "user":
                _t += 1
            _turn_of[m["id"]] = _t
        for p in pulled:
            if used >= pull_cap or not p["content"]:
                break
            c = _clip(p["content"], max(pull_cap - used, 16))
            if not c:
                continue
            head = "【相关历史片段（第 %d 轮 · %s，相关度 %s）】" % (
                _turn_of.get(p["id"], 0), "用户" if p.get("role") == "user" else "助手", p["score"])
            pull_blocks.append({"role": "system", "content": head + "\n" + c})
            used += _ct(c) + _ct(head)
        # ③ 其他话题分话题摘要（剩余预算，低优先级不足则停；排除当前话题组避免重复）
        for g in groups:
            if g is cur_group:
                continue
            if used >= budget_tok:
                break
            seg = [m for m in g["msgs"] if m["id"] not in used_ids]
            if not seg:
                continue
            s = self._history_summary(conversation_id, f"topic:{g['topic'][:40]}", seg, summary_chars)
            if not s:
                continue
            c = _clip(s, max(0, budget_tok - used))
            if not c:
                break
            sum_blocks.append({"role": "system", "content": f"【话题：{g['topic']} 摘要】\n{c}"})
            used += _ct(c)
        # 输出按时间序：会话产物摘要（最稳定的事实）→ 话题摘要（最老）→ 拉回（中间）
        # → 当前话题原文（最近，紧邻当前输入）
        return _digest + sum_blocks + pull_blocks + raw_list

    def _tag_topics(self, conn, conversation_id):
        """会话消息话题打标（惰性全量重算+回写 messages.topic，规则确定性）。
        相邻用户消息与「当前话题段代表向量」（段内全部消息 bigram 累加）余弦 < 阈值 → 话题切换
        （以承接词开头则继承当前话题）；助手消息继承前一条用户消息话题并并入段代表向量。
        返回 {message_id: topic}。
        """
        from collections import Counter
        from core import config as _cfg
        from knowledge_engine import VectorEngine
        threshold = float(_cfg.get("context", "topic_sim_threshold", 0.15))
        label_chars = int(_cfg.get("context", "topic_label_chars", 14))
        rows = conn.execute(
            "SELECT id, role, content, topic FROM messages WHERE conversation_id=? AND role IN ('user','assistant') "
            "ORDER BY id", (conversation_id,)).fetchall()
        if not rows:
            return {}
        mapping = {r["id"]: (r["topic"] or "") for r in rows}
        # 无未打标消息 → 直接复用
        if all(t for t in mapping.values()):
            return mapping
        ve = VectorEngine()
        # 延续词：承接式追问继承当前话题；回指词（刚才/前面/之前）不在此列，
        # 其与段代表向量相似度低时自然开新话题 → 语义拉回负责把回指的历史原文拉回
        # 2026-09-29（用户五轮反馈1）：补"重试/重跑/重新/再来"——任务终止/失败后用户打"重试"，
        # 若按新话题切走，当前话题原文就只剩"重试"两个字，AI 表现为"不知道之前做了什么"。
        carry = TOPIC_CARRY_WORDS   # P1-4：单一来源，见 common.py（与 _is_topic_switch 共用）
        cur_topic = ""
        topic_vec = None  # 段代表向量（段内消息 bigram 累加，随对话增长）
        last_user_id = None
        for r in rows:
            content = (r["content"] or "").strip()
            if r["role"] == "user":
                vec = ve._vector(content) if content else None
                if topic_vec is None:
                    cur_topic = content[:label_chars] or "话题1"
                    topic_vec = Counter(vec) if vec else Counter()
                elif vec is not None:
                    sim = ve._cosine(vec, topic_vec)
                    # 切换需「相似度低 + 无共同 bigram 词」双重条件：短承接句（仅 1 个共享业务词）
                    # 相似度天然低但语义承接（如"链路余量校核"承接"链路预算"），默认偏不切，
                    # 宁多合并（靠语义拉回兜底）也不误切
                    shared = bool(set(vec.keys()) & set(topic_vec.keys()))
                    if is_topic_switch(sim, shared, content.startswith(carry), threshold):
                        cur_topic = content[:label_chars] or cur_topic
                        topic_vec = Counter(vec)  # 新话题段
                    else:
                        topic_vec.update(vec)  # 承接：并入段代表
                last_user_id = r["id"]
            else:  # assistant 继承前一条用户消息话题，并入段代表
                cur_topic = mapping.get(last_user_id, cur_topic) if last_user_id is not None else cur_topic
                if topic_vec is not None:
                    topic_vec.update(ve._vector(content))
            mapping[r["id"]] = cur_topic
            conn.execute("UPDATE messages SET topic=? WHERE id=?", (cur_topic, r["id"]))
        conn.commit()
        return mapping

    def _is_topic_switch(self, conversation_id, user_input, _max_msgs: int = 40) -> bool:
        """P1-4 L2：本轮输入是否**换了话题**（决定 `_merge_slots` 是否丢弃历史槽位）。

        取段：本轮消息尚未落库（execute/stream 是**先回复后落库**），无法走 `_tag_topics`
        的全量循环 → 取**最后一个已打标话题段**作为当前话题。

        ⚠️ 判据**不直接复用** `_tag_topics` 的两个条件，实测理由（都是踩过才写的）：
          ① `_tag_topics` 的「段代表向量」是**全段 user+assistant 累加**且随对话无限增长
             —— 实测 conv=514 达 7879 个 bigram，于是「无共同 bigram」条件对**任何**输入
             恒为假（异话题句也命中 '分析'/'报告'/'项目' 等泛词）→ 判据永不触发。
             佐证：该会话全程只被分成 **1 个**话题段（不同话题的消息混在同段）。
          ② 复用 `context.topic_sim_threshold=0.15` 也不行：该阈值是在**大词袋**下标定的，
             同话题句被稀释到 0.09~0.19 → 实测**误切 4/7**（把追问轮的对象清掉）。
        故改为：**只累加段内 user 消息**（词袋 ~310，不随 assistant 长文膨胀）+
        独立键 `slots.topic_switch_threshold`。标定见 `tmp/mt_ctx/calib2_variants.py`
        （conv=514 真实段 + 8 同话题 / 8 异话题样本，四种构造对比）:

            构造         同话题区间           异话题区间           间隙比  词袋
            A 全段累加   [0.0922, 0.1889]   [0.0000, 0.0475]   1.94x  7879
            C user 累加  [0.1342, 0.4813]   [0.0000, 0.0350]   3.83x   310   ← 采用

        取 0.08（落在 C 的分离带 (0.0350, 0.1342) 内）：8/8 同话题不切、8/8 异话题切。

        任何异常 / 无历史 / 未打标 → False（**保守：不重置**）。
        宁可多留（有上限兜底），也不因判据抖动把追问轮的对象清掉。
        """
        try:
            if not conversation_id or int(conversation_id) <= 0:
                return False
            txt = (user_input or "").strip()
            if not txt:
                return False
            from core import config as _cfg
            if not bool(_cfg.get("slots", "reset_on_topic_switch", True)):
                return False
            # 独立于 `context.topic_sim_threshold`（0.15）—— 那个阈值对应「全段累加」的
            # 大词袋；本判据用「段内 user 消息累加」，标定见方法 docstring。
            threshold = float(_cfg.get("slots", "topic_switch_threshold", 0.08))
            _ref_n = int(_cfg.get("slots", "topic_ref_msgs", 10) or 10)
            conn = get_db()
            try:
                rows = conn.execute(
                    "SELECT id, role, content, topic FROM messages "
                    "WHERE conversation_id=? AND role IN ('user','assistant') "
                    "ORDER BY id DESC LIMIT ?",
                    (int(conversation_id), int(_max_msgs))).fetchall()
            finally:
                conn.close()
            rows = list(reversed(rows))
            if not rows:
                return False
            # ⚠️ 尾部的消息**可能尚未打标** —— `_tag_topics` 是**惰性全量重算**（全部已打标
            #   才直接复用），而落库与打标之间存在窗口：实测 conv=514 最后 4 条
            #   （3010~3013，正是最近一轮）topic 为空。若直接取 `rows[-1]["topic"]`，
            #   判据在这些会话上**恒为 False**（换了话题也不重置），整个 L2 形同未接。
            #   故取**最后一条已打标**消息的 topic 作为当前话题，其后的未打标消息
            #   并入该段（它们是当前话题的延续，内容同样计入代表向量）。
            k = None
            for i in range(len(rows) - 1, -1, -1):
                if (rows[i]["topic"] or ""):
                    k = i
                    break
            if k is None:
                return False   # 整段都没打过标 → 保守不重置
            last_topic = rows[k]["topic"]
            # 只取 **user** 消息：话题由用户说的话定义；assistant 的长篇产出混进来会把
            # 词袋从 ~300 撑到 ~8000，既稀释相似度又让「无共同 bigram」恒真（见 docstring）。
            seg = [rows[i]["content"] or "" for i in range(k + 1, len(rows))
                   if rows[i]["role"] == "user"]
            for i in range(k, -1, -1):
                if (rows[i]["topic"] or "") != last_topic:
                    break
                if rows[i]["role"] == "user":
                    seg.append(rows[i]["content"] or "")
            seg.reverse()
            seg = seg[-_ref_n:]   # 段内最近 N 条 user 消息（词袋规模稳定）
            from collections import Counter
            ve = VectorEngine()
            topic_vec = Counter()
            for c in seg:
                topic_vec.update(ve._vector(c))
            if not topic_vec:
                return False
            vec = ve._vector(txt)
            if not vec:
                return False
            sim = ve._cosine(vec, topic_vec)
            # use_shared=False：标定集上「仅相似度」已 8/8 + 8/8 完全分离；再加 shared 条件
            # 只会带来漏切风险（`_tag_topics` 加它是因为它的词袋大到 shared 恒真、需要额外约束，
            # 此处词袋小、无此问题）。
            return bool(is_topic_switch(sim, True, txt.startswith(TOPIC_CARRY_WORDS),
                                        threshold, use_shared=False))
        except Exception:
            return False

    def _semantic_scores(self, texts, query):
        """统一语义出口：真 embedding（texts+query 同批向量化保维度一致）→ bigram 降级。

        返回 (scores, backend)：backend='dense' 时 scores 与 texts 对齐（embedding 余弦 0~1）；
        否则 (None, 'bigram')，调用方走 VectorEngine bigram 打分。
        embedding.enabled=False（Mock/离线确定性）强制 bigram；Embedder 自带 probe/query 缓存。
        """
        try:
            from core import config as _cfg_s
            if not _cfg_s.as_bool("embedding", "enabled", True):
                return None, "bigram"
            if not query or not texts:
                return None, "bigram"
            from knowledge_pipeline import Embedder
            from database import get_db
            conn = get_db()
            try:
                ed = Embedder(conn)
            finally:
                conn.close()
            all_texts = list(texts) + [query]
            vecs, version = ed.embed_with_version(all_texts, batch_size=0)
            if version != "openai-compat" or not vecs or len(vecs) != len(all_texts):
                return None, "bigram"
            import math as _math
            qv = vecs[-1]
            qn = _math.sqrt(sum(x * x for x in qv)) or 1.0
            scores = []
            for v in vecs[:-1]:
                if not v or len(v) != len(qv):
                    scores.append(0.0)
                    continue
                vn = _math.sqrt(sum(x * x for x in v)) or 1.0
                scores.append(sum(a * b for a, b in zip(qv, v)) / (qn * vn))
            return scores, "dense"
        except Exception:
            return None, "bigram"

    def _search_history(self, msgs, query, topk=6, threshold=0.15, cur_topic="", per_role_cap=0):
        """会话内语义拉回：真 embedding 优先（同批向量化余弦），bigram 降级。

        dense 路：embedding 余弦独立量纲，阈值用 topic_retrieve_threshold_dense；
        话题域加权不再需要（稠密向量对措辞差异鲁棒，query 已拼当前话题文本）。
        bigram 路：余弦打分 + 话题域加权（候选 topic 与当前话题 ts≥0.2 加 0.25*ts，
        补偿细节词在长 query 中被稀释）；候选截最近 120 条控 embedding 成本。

        `per_role_cap`（P1-8，默认 0=不限制，回调方传 3）：**输入/输出分角色配额**。
        为什么必须分：短文本在余弦打分上系统性偏高（模长归一化后语义方向更"纯"）——
        实测（真库）13 字符的追问得 **0.988**、4 字符的「继续优化」得 0.66，而助手历轮的
        数千字结论只有 0.5~0.6。修好 `used_ids` 后 conv=1 的混合 top-6 **全是 user**，
        助手此前产出的结论/版本一条都进不来 —— 那既不是「相关性评估」也不符合
        「输入、输出都应被评估」的要求。配额后：先各取 ≤cap，**名额未满再按分数补齐**
        （单角色场景不因配额减产）。返回项带 `role`，调用方据此拼可追溯块头。
        """
        from core import config as _cfg
        from core import config as _cfg_s
        cands = [(m, (m.get("content") or "")[:400]) for m in msgs if (m.get("content") or "")[:400]]
        if not query or not cands:
            return []
        cands = cands[-120:]  # 成本护栏：拉回只需 top-k，候选过多按时间取最近
        dense_th = float(_cfg_s.get("context", "topic_retrieve_threshold_dense", 0.35))
        scores, backend = self._semantic_scores([c for _, c in cands], query)
        if backend == "dense":
            scored = [(s, m) for (m, _), s in zip(cands, scores) if s >= dense_th]
        else:
            from knowledge_engine import VectorEngine
            ve = VectorEngine()
            qv = ve._vector(query or "")
            if not qv:
                return []
            tqv = ve._vector(cur_topic) if cur_topic else None
            scored = []
            for m, c in cands:
                s = ve._cosine(qv, ve._vector(c))
                if tqv:
                    mt = ve._vector(m.get("topic") or "")
                    ts = ve._cosine(tqv, mt) if mt else 0.0
                    if ts >= 0.2:
                        s += 0.25 * ts  # 同话题域加权（上限约 +0.25）
                if s >= threshold:
                    scored.append((s, m))
        scored.sort(key=lambda x: x[0], reverse=True)
        picked = []
        if per_role_cap and per_role_cap > 0:
            cnt = {}
            for s, m in scored:                     # 第一轮：分角色配额
                if len(picked) >= topk:
                    break
                r = m.get("role") or ""
                if cnt.get(r, 0) >= per_role_cap:
                    continue
                cnt[r] = cnt.get(r, 0) + 1
                picked.append((s, m))
            if len(picked) < topk:                  # 第二轮：名额没满 → 按分数补齐（防单角色减产）
                chosen = {m["id"] for _s, m in picked}
                for s, m in scored:
                    if len(picked) >= topk:
                        break
                    if m["id"] in chosen:
                        continue
                    picked.append((s, m))
        else:
            picked = scored[:topk]
        return [{"id": m["id"], "role": m["role"], "content": m["content"][:300], "score": round(s, 3)}
                for s, m in picked]

    # ── 会话历史注入 v1（旧三层窗口，话题打标失败时兜底）──
    def _load_history_legacy(self, conversation_id, limit=None):
        from core import config as _cfg
        limit = limit or int(_cfg.get("context", "history_immediate_turns", 6))
        mid_turns = int(_cfg.get("context", "history_mid_turns", 20))
        msg_max = int(_cfg.get("context", "history_msg_max_chars", 1500))
        summary_chars = int(_cfg.get("context", "history_summary_chars", 800))
        try:
            conn = get_db()
            rows = conn.execute(
                "SELECT id, role, content FROM messages WHERE conversation_id=? AND role IN ('user','assistant') "
                "ORDER BY id", (conversation_id,)).fetchall()
            conn.close()
        except Exception:
            return []
        if not rows:
            return []
        msgs = []
        for r in rows:
            c = (r["content"] or "")[:msg_max]
            msgs.append({"id": r["id"], "role": r["role"], "content": c})
        immediate = msgs[-(limit * 2):] if limit else msgs
        older = msgs[:-(limit * 2)] if (limit and len(msgs) > limit * 2) else []
        out = []
        if older:
            mid_seg = older[-(mid_turns * 2):]
            hist_seg = older[:-(mid_turns * 2)] if len(older) > mid_turns * 2 else []
            mid_sum = self._history_summary(conversation_id, "mid", mid_seg, summary_chars)
            if mid_sum:
                out.append({"role": "system", "content": f"【更早对话摘要】\n{mid_sum}"})
            if hist_seg:
                hist_sum = self._history_summary(conversation_id, "hist", hist_seg, summary_chars)
                if hist_sum:
                    out.append({"role": "system", "content": f"【历史对话概览】\n{hist_sum}"})
        for m in immediate:
            out.append({"role": m["role"], "content": m["content"]})
        return out

    def _history_summary(self, conversation_id, scope, seg, max_chars=800):
        """生成/复用某段历史的摘要（缓存锚点 = 段内最大消息 id，段变化自动重算）。"""
        if not seg:
            return ""
        anchor = seg[-1]["id"]
        try:
            conn = get_db()
            row = conn.execute(
                "SELECT summary FROM conversation_summaries WHERE conversation_id=? AND scope=? AND anchor_id=?",
                (conversation_id, scope, anchor)).fetchone()
            if row and row["summary"]:
                return row["summary"][:max_chars]
            text = "\n".join(f"{'用户' if m['role'] == 'user' else '助手'}: {m['content'][:300]}" for m in seg)
            summary = self._llm_summarize(text, scope)
            conn.execute(
                "INSERT OR REPLACE INTO conversation_summaries (conversation_id, scope, anchor_id, summary) "
                "VALUES (?,?,?,?)",
                (conversation_id, scope, anchor, summary))
            conn.commit()
            conn.close()
            return summary[:max_chars]
        except Exception:
            return ""

    def _llm_summarize(self, text, scope):
        """LLM 摘要（保留决策/实体/待办，丢弃客套）；Mock/失败 → 规则降级（首句拼接）。"""
        try:
            from llm import llm_client
            prompt = ("将以下对话历史压缩为简洁摘要：保留关键决策、实体、约束与待办；"
                      "丢弃客套与重复内容。输出 200 字以内。\n\n对话历史：\n" + text[:4000])
            resp = llm_client.chat([{"role": "user", "content": prompt}], _intent="history_summary")
            s = ((resp.get("choices") or [{}])[0].get("message", {}) or {}).get("content") or ""
            s = s.strip()
            if len(s) > 30:
                return s
        except Exception:
            pass
        parts = []
        for line in text.split("\n"):
            if line.startswith(("用户:", "助手:")) and line[3:].strip():
                parts.append(line[:80])
        return "；".join(parts[:10]) or ""

    @staticmethod
    def _truncate_budget(text: str, max_chars: int, keep_head: bool = True) -> str:
        """P1a-2 Token 预算裁剪（字符版兜底）：超限时保留头部（检索数据相关性在头部，对齐 Lost-in-the-Middle 经验）。

        粗估：1 中文字 ≈ 1.5 token，字符上限 ≈ token 预算 × 2/3。
        """
        if not text or max_chars <= 0 or len(text) <= max_chars:
            return text
        if keep_head:
            return text[:max_chars] + "\n…（超预算已裁剪）"
        return "…（超预算已裁剪）\n" + text[-max_chars:]

    @staticmethod
    def _truncate_tokens(text: str, max_tokens: int, keep_head: bool = True, model_hint: str | None = None) -> str:
        """T6 Token 驱动裁剪：按 token 估算超限后以字符近似比例截断并复测（最多 2 次迭代，防中英混排误差）。

        返回裁剪后文本；超限时附裁剪标记。未超限原样返回。
        """
        from core.token_counter import count_tokens
        if not text or max_tokens <= 0:
            return text
        n = count_tokens(text, model_hint)
        if n <= max_tokens:
            return text
        # 裁剪标记自身也占预算：预留标记 token 开销，避免"裁到上限+标记"最终超限
        marker = "\n…（超预算已裁剪）" if keep_head else "…（超预算已裁剪）\n"
        budget = max_tokens - count_tokens(marker, model_hint)
        if budget < 16:
            budget = max_tokens  # 预算过小时退回原值（标记开销可忽略）
        # 粗比例：字符数 / token 数 → 目标字符数；逐轮逼近（token 数非严格线性）
        ratio = max(len(text) / max(n, 1), 0.5)
        target_chars = int(budget * ratio)
        for _ in range(2):
            cut = text[:target_chars] if keep_head else text[-target_chars:]
            cn = count_tokens(cut, model_hint)
            if cn <= budget:
                return (cut + marker) if keep_head else (marker + cut)
            target_chars = int(target_chars * budget / max(cn, 1))
        cut = text[:target_chars] if keep_head else text[-target_chars:]
        return (cut + marker) if keep_head else (marker + cut)

    def _apply_context_budget(self, system_prompt: str, retrieval_text: str, history_len: int,
                             scale: float = 1.0) -> str:
        """P1a-2/T6 上下文预算分区：检索区与历史区分别裁剪，防超窗（对齐 Levelop Context Spec）。

        T6 升级：以 token 预算为准（budget_retrieval_tokens / budget_history_tokens），
        旧字符配置（budget_retrieval_chars 等）保留为估算回退路径兜底。
        P1-4b：占比制优先（context.budget_{retrieval,history}_ratio > 0 时按窗口比例算）。
        返回裁剪后的 system_prompt。检索区保留头部（最相关），历史区保留尾部（最近）。
        """
        try:
            from core import config as _cfg
            from core.token_counter import count_tokens
            # P1-2：scale 是**总闸降档系数**（1.0=默认行为；<1 时检索/历史同步收紧）。
            # 只由 _apply_total_budget 传入；两个既有调用点不传 ⇒ 行为不变。
            retr_tok = int(ctx_budget_tokens("retrieval", "budget_retrieval_chars", 4000) * scale)
            hist_tok = int(ctx_budget_tokens("history", "budget_history_chars", 3000) * scale)
            if retrieval_text:
                system_prompt = system_prompt.replace(
                    retrieval_text,
                    self._truncate_tokens(retrieval_text, retr_tok, keep_head=True),
                    1)
            # 历史区：system_prompt 中「更早对话摘要/历史对话概览」两块截断（token 驱动）
            def _cut_hist(m, half):
                seg = m.group(2)
                if count_tokens(seg) <= (hist_tok // 2 if half else hist_tok):
                    return m.group(1) + seg
                return m.group(1) + self._truncate_tokens(seg, hist_tok // 2 if half else hist_tok, keep_head=True)
            system_prompt = re.sub(
                r"(【更早对话摘要】\n)(.{1,8000}?)(?=\n【历史对话概览】|$)",
                lambda m: _cut_hist(m, True), system_prompt, flags=re.S)
            system_prompt = re.sub(
                r"(【历史对话概览】\n)(.{1,8000}?)(?=\n【|$)",
                lambda m: _cut_hist(m, False), system_prompt, flags=re.S)
            return system_prompt
        except Exception:
            return system_prompt

    def _apply_total_budget(self, messages, context_text: str):
        """P1-2（2026-09-24）上下文**总闸**：分段预算之和可超窗，发送前算总账、超了降档重裁。

        背景：检索 4k + 历史 3k + system 本体 4~5k + 附件 2.7k + v2 约束 1.3k ≈ 12~16k，
        叠加超窗时由输出侧 _ctx_guard 把 max_tokens 夹小 ⇒ 输出被悄悄截短
        （即 S4 登记过的「报告在结论前中断」——输入被切与输出被夹，现象难分）。

        · 常态（预算内）**零开销**：只算一次 token 就返回
        · 降档 = 用**同一套** _apply_context_budget 换更紧的 scale 重跑（不引入第二套裁剪逻辑）
        · system 本体（建模规则）不在降档范围——压它会直接改变产出质量
        · 可关：context.total_budget_guard=false；上限 context.total_budget_tokens（默认 10000）
        · 每次降档打日志留痕（before/after/档位）——总闸若静默，等于没做
        """
        try:
            from core import config as _cfg
            from core.token_counter import count_messages_tokens
            if not messages:
                return messages
            if not _cfg.as_bool("context", "total_budget_guard", True):
                return messages
            total_cap = int(_cfg.get("context", "total_budget_tokens", 10000) or 10000)
            before = count_messages_tokens(messages)
            if before <= total_cap:
                return messages
            sys0 = messages[0].get("content") or ""
            after = before
            for step, scale in enumerate((0.6, 0.35), start=1):
                sys0 = self._apply_context_budget(sys0, context_text or "", len(messages), scale=scale)
                messages[0]["content"] = sys0
                after = count_messages_tokens(messages)
                if after <= total_cap:
                    break
            try:
                import logging as _lg
                _lg.getLogger("mbse.llm").warning(
                    "[context_total] 总闸触发：估算输入 %d > 预算 %d，降档后 %d "
                    "（裁掉 %d tok，到达档位 %d，messages=%d 条）",
                    before, total_cap, after, before - after, step, len(messages))
            except Exception:
                pass
            return messages
        except Exception:  # noqa: BLE001 —— 总闸自身故障绝不阻断主链路，但必须留痕
            import traceback as _tb
            _tb.print_exc()
            return messages

    # ── 工作流匹配：词法 + 语义双通道 + RRF 融合（行业对齐：Voiceflow 混合检索 / juejin RRF / 百度漏斗式）──

    # ── 续作短语判别（2026-09-29 用户五轮反馈1）──
    #  场景：任务终止/失败后用户只打"重试"/"继续"——这不是新任务，是要求 AI 衔接上文。
    #  此前两个断点叠加导致"AI 不知道之前的内容"：
    #    ① carry 承接词表没有"重试"→ 话题被打成新段，当前话题原文只剩"重试"两字；
    #    ② 被终止的那轮产出根本没落库（见 stream.py 尾部 except 固化逻辑）→ 历史里无货可接。
    #  本判别供 stream.py 在组装 prompt 时追加显式"续作指令"。
    _CONTINUATION_RE = None

    @classmethod
    def _is_continuation_input(cls, text):
        """是否为纯续作短语（重试/继续/接着来/retry…）。判据：短（≤12字）且命中模式。

        2026-09-29 修订：初版漏了「重试一下 / 重跑一遍 / 继续一次」这类**量词后缀**
        （自检 L1 正例 '重试一下' 未命中 → 会被判成新任务）。现把「一下吧/一遍/一回/一次」
        收进统一后缀组，并对每个动词后统一允许。
        """
        import re as _re
        if cls._CONTINUATION_RE is None:
            # 量词后缀（可省）：一下/一遍/一回/一次/一个/下/遍
            _suf = r"(?:一下|一遍|一回|一次|一个|下|遍)?"
            cls._CONTINUATION_RE = _re.compile(
                r"^(?:"
                r"重试" + _suf + r"|"
                r"重跑" + _suf + r"|"
                r"重新" + _suf + r"(?:来|跑|执行|生成)" + _suf + r"|"
                r"再来" + _suf + r"|"
                r"继续" + _suf + r"(?:吧|执行|做)?" + _suf + r"|"
                r"接着" + _suf + r"(?:来|做|干)?" + _suf + r"|"
                r"go\s*on|retry|continue"
                r")[吧，,。！!？?～~\s]*$",
                _re.IGNORECASE)
        t = (text or "").strip()
        return bool(t) and len(t) <= 12 and bool(cls._CONTINUATION_RE.match(t))
