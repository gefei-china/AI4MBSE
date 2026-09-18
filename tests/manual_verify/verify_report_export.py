"""报告生成技能 + 多格式文档导出（md/docx/pdf）闭环验证。

覆盖：
T1.  export_markdown：Report → 完整 markdown 文本
T2.  export_docx：Word 文档生成（python-docx 读回：标题/分节/表格）
T3.  export_pdf：PDF 生成（%PDF 文件头 + 体积）
T4.  export() 统一出口三格式（bytes）
T5.  **粗体** 富文本 → docx 加粗 run
T6.  markdown 表格 → docx 表格单元格
T7.  API POST /api/report/export：md/docx/pdf 均 200 + Content-Type + Content-Disposition
T8.  API 非法格式 → 400
T9.  技能「报告生成」注册 published + allowed_tools/references/scripts
T10. 报告生成技能脚本 run(fmt=docx/pdf/md) 产出文件
T11. 行业调研技能脚本 run(fmt=docx/pdf) 产出文件
"""
import importlib.util
import json
import os
import shutil
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
os.environ["MBSE_LLM_FORCE_MOCK"] = "1"
_TEST_DB = os.path.join(os.path.dirname(__file__), "_report_export_test.db")
if os.path.exists(_TEST_DB):
    os.remove(_TEST_DB)
os.environ["MBSE_DB_PATH"] = _TEST_DB

from database import db_conn, init_db  # noqa: E402

init_db()
import register_skills  # noqa: E402
register_skills.main()

from report_generator import report_generator  # noqa: E402

PASS = 0
FAIL = 0


def chk(name, ok, detail=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [PASS] {name}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name}  → {detail}")


WORK = os.path.join(os.path.dirname(__file__), "_re_work")
if os.path.exists(WORK):
    shutil.rmtree(WORK)
os.makedirs(WORK)

REPORT = {
    "title": "宽带通信卫星载荷分析报告",
    "sections": [
        {"heading": "概述", "body": "本报告分析**卫星载荷**的关键指标，覆盖 5 个型号。\n\n"
         "| 型号 | 功率 | 带宽 |\n|---|---|---|\n| A | 5kW | 2GHz |\n| B | 8kW | 4GHz |\n\n"
         "- 指标一\n- 指标二"},
        {"heading": "问题与风险", "body": "风险清单如下：\n1. 散热不足\n2. 频率冲突"},
        {"heading": "结论", "body": "**结论**：A 型适配当前需求。"},
    ],
    "summary": "结论：A 型适配。",
    "report_type": "analysis",
}

print("== T1: export_markdown ==")
md = report_generator.export_markdown(REPORT)
chk("T1 md 含标题", f"# {REPORT['title']}" in md, md[:80])
chk("T1 md 含分节与表格", "## 概述" in md and "| A | 5kW | 2GHz |" in md, md[:200])

print("== T2: export_docx ==")
dpath = os.path.join(WORK, "报告.docx")
r = report_generator.export_docx(REPORT, dpath)
chk("T2 docx 落盘", r == dpath and os.path.isfile(dpath), str(r))
from docx import Document
doc = Document(dpath)
paras = [p.text for p in doc.paragraphs]
chk("T2 docx 标题", paras and paras[0] == REPORT["title"], str(paras[:2]))
chk("T2 docx 分节 heading", "概述" in paras and "问题与风险" in paras and "结论" in paras, str(paras[:8]))
chk("T2 docx 表格存在", len(doc.tables) >= 1, f"tables={len(doc.tables)}")

print("== T3: export_pdf ==")
ppath = os.path.join(WORK, "报告.pdf")
r = report_generator.export_pdf(REPORT, ppath)
pdf_head = open(ppath, "rb").read(5)
chk("T3 pdf 落盘且 %PDF 头", r == ppath and os.path.isfile(ppath) and pdf_head == b"%PDF-", str(pdf_head))
chk("T3 pdf 体积 > 1KB", os.path.getsize(ppath) > 1024, os.path.getsize(ppath))

print("== T4: export() 统一出口 ==")
b_md = report_generator.export(REPORT, "md")
b_docx = report_generator.export(REPORT, "docx")
b_pdf = report_generator.export(REPORT, "pdf")
chk("T4 md bytes", isinstance(b_md, bytes) and REPORT["title"] in b_md.decode("utf-8"))
chk("T4 docx bytes", isinstance(b_docx, bytes) and b_docx[:2] == b"PK", b_docx[:4])
chk("T4 pdf bytes", isinstance(b_pdf, bytes) and b_pdf[:5] == b"%PDF-", b_pdf[:5])
try:
    report_generator.export(REPORT, "xlsx")
    chk("T4 非法格式抛错", False, "no raise")
except ValueError:
    chk("T4 非法格式抛错", True)

print("== T5: **粗体** → docx 加粗 run ==")
bold_texts = [run.text.strip() for p in doc.paragraphs for run in p.runs if run.bold and run.text.strip()]
chk("T5 存在加粗 run（卫星载荷/结论）",
    any("卫星载荷" in t for t in bold_texts) and any("结论" in t for t in bold_texts),
    str(bold_texts)[:120])

print("== T6: markdown 表格 → docx 表格 ==")
tb = doc.tables[0]
hdr = [c.text for c in tb.rows[0].cells]
chk("T6 表头 型号/功率/带宽", hdr[:3] == ["型号", "功率", "带宽"], str(hdr))
chk("T6 数据行 A/5kW/2GHz", [c.text for c in tb.rows[1].cells] == ["A", "5kW", "2GHz"],
    str([c.text for c in tb.rows[1].cells]))

print("== T7: API 导出 ==")
from fastapi.testclient import TestClient  # noqa: E402
from main import app  # noqa: E402
client = TestClient(app)
for fmt, ctype, magic in (("md", "text/markdown", None),
                          ("docx", "application/vnd.openxmlformats", b"PK"),
                          ("pdf", "application/pdf", b"%PDF-")):
    resp = client.post("/api/report/export", json={**REPORT, "fmt": fmt})
    chk(f"T7 {fmt} 200 + Content-Type", resp.status_code == 200 and ctype in resp.headers.get("content-type", ""),
        f"status={resp.status_code} ct={resp.headers.get('content-type')}")
    chk(f"T7 {fmt} Content-Disposition 附件", "attachment" in resp.headers.get("content-disposition", ""),
        resp.headers.get("content-disposition"))
    if magic:
        chk(f"T7 {fmt} 文件魔数", resp.content[:len(magic)] == magic, resp.content[:6])
    else:
        chk(f"T7 {fmt} 内容含标题", REPORT["title"] in resp.content.decode("utf-8", "replace"))
resp = client.post("/api/report/export", json={**REPORT, "fmt": "xlsx"})
chk("T8 非法格式 400", resp.status_code == 400, resp.status_code)

print("== T9: 报告生成技能注册 ==")
with db_conn() as conn:
    row = conn.execute("SELECT name, skill_type, status, allowed_tools, `references`, scripts FROM skills WHERE name='报告生成'").fetchone()
chk("T9 报告生成 published+package", row and row["status"] == "published" and row["skill_type"] == "package", dict(row) if row else "missing")
if row:
    at = json.loads(row["allowed_tools"] or "[]")
    chk("T9 allowed_tools 含 graph_retrieve+file_write", "graph_retrieve" in at and "file_write" in at, str(at))
    chk("T9 references/scripts 清单", len(json.loads(row["references"] or "[]")) >= 1
        and len(json.loads(row["scripts"] or "[]")) >= 1, f"refs={row['references']} scripts={row['scripts']}")

print("== T10: 报告生成技能脚本 ==")
with db_conn() as conn:
    pkg = conn.execute("SELECT package_path FROM skills WHERE name='报告生成'").fetchone()["package_path"]
spec = importlib.util.spec_from_file_location("_skill_report_gen", os.path.join(pkg, "scripts", "main.py"))
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)
for fmt in ("md", "docx", "pdf"):
    out_path = os.path.join(WORK, f"技能报告.{fmt}")
    out = mod.run("技能生成测试报告", markdown="## 概述\n这是**正文**内容。\n\n| 项 | 值 |\n|---|---|\n| A | 1 |", fmt=fmt, 输出路径=out_path)
    chk(f"T10 run fmt={fmt} 产出文件", os.path.isfile(out_path) and fmt.upper() in str(out), str(out)[:200])

print("== T11: 行业调研技能脚本 ==")
with db_conn() as conn:
    pkg2 = conn.execute("SELECT package_path FROM skills WHERE name='行业调研'").fetchone()["package_path"]
spec2 = importlib.util.spec_from_file_location("_skill_industry", os.path.join(pkg2, "scripts", "main.py"))
mod2 = importlib.util.module_from_spec(spec2)
spec2.loader.exec_module(mod2)
for fmt in ("docx", "pdf"):
    out_path = os.path.join(WORK, f"行业调研报告.{fmt}")
    out = mod2.run("AI Agent 平台搭建", 输出路径=out_path, fmt=fmt)
    chk(f"T11 行业调研 fmt={fmt} 产出文件", os.path.isfile(out_path) and fmt.upper() in str(out), str(out)[:200])
out_md = os.path.join(WORK, "行业调研报告.md")
mod2.run("AI Agent 平台搭建", 输出路径=out_md, fmt="md")
md_txt = open(out_md, encoding="utf-8").read()
chk("T11 行业调研 md 含六模块", "Tools" in md_txt and "MCP" in md_txt and "多Agent" in md_txt, md_txt[:100])

shutil.rmtree(WORK, ignore_errors=True)
if os.path.exists(_TEST_DB):
    os.remove(_TEST_DB)
print(f"\n===== 结果: {PASS} 通过 / {FAIL} 失败 =====")
sys.exit(1 if FAIL else 0)
